from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path(
    "docs/partnerships/Fomo-Live-Feed-Frontrun-Cooperation-Proposal.zh-CN.docx"
)

FONT_CN = "STSong"
FONT_LATIN = "Arial"
INK = "17202A"
MUTED = "5C6773"
BLUE = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
GRID = "D9D9D9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = p_pr.find(qn("w:keepNext"))
    if value and keep_next is None:
        keep_next = OxmlElement("w:keepNext")
        p_pr.append(keep_next)
    elif not value and keep_next is not None:
        p_pr.remove(keep_next)


def set_keep_together(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep_lines = p_pr.find(qn("w:keepLines"))
    if value and keep_lines is None:
        keep_lines = OxmlElement("w:keepLines")
        p_pr.append(keep_lines)
    elif not value and keep_lines is not None:
        p_pr.remove(keep_lines)


def set_run_font(run, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT_CN)
    r_pr = run._element.get_or_add_rPr()
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def style_paragraph_runs(paragraph, size=11, color=INK) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, color=color)


def add_rich_paragraph(doc, parts, *, style=None, before=0, after=6, keep=True):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.line_spacing = 1.2
    if keep:
        set_keep_together(paragraph)
    for text, bold in parts:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, bold=bold, color=INK)
    return paragraph


def add_bullet(doc, text: str, level=0):
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    set_keep_together(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.6, color=INK)
    return paragraph


def add_number(doc, number: int, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.24)
    paragraph.paragraph_format.first_line_indent = Inches(-0.24)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    set_keep_together(paragraph)
    run = paragraph.add_run(f"{number}.  {text}")
    set_run_font(run, size=10.6, color=INK)
    return paragraph


def add_heading(doc, text: str, level=1):
    paragraph = doc.add_heading(text, level=level)
    set_keep_with_next(paragraph)
    return paragraph


def add_table(doc, headers, rows, widths, *, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, text, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        set_run_font(run, size=font_size, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        body_row = table.add_row()
        tr_pr = body_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        cells = body_row.cells
        for col_index, (cell, text, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.1
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(text))
            set_run_font(run, size=font_size, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal_lang = normal._element.rPr.find(qn("w:lang"))
    if normal_lang is None:
        normal_lang = OxmlElement("w:lang")
        normal._element.rPr.append(normal_lang)
    normal_lang.set(qn("w:val"), "en-US")
    normal_lang.set(qn("w:eastAsia"), "zh-CN")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    title = doc.styles["Title"]
    title.font.name = FONT_LATIN
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title_lang = title._element.rPr.find(qn("w:lang"))
    if title_lang is None:
        title_lang = OxmlElement("w:lang")
        title._element.rPr.append(title_lang)
    title_lang.set(qn("w:val"), "en-US")
    title_lang.set(qn("w:eastAsia"), "zh-CN")
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 15, 12, 6),
        ("Heading 2", 12, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style_lang = style._element.rPr.find(qn("w:lang"))
        if style_lang is None:
            style_lang = OxmlElement("w:lang")
            style._element.rPr.append(style_lang)
        style_lang.set(qn("w:val"), "en-US")
        style_lang.set(qn("w:eastAsia"), "zh-CN")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "Fomo Live Feed  ×  Frontrun"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    style_paragraph_runs(header, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def build_document() -> Document:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Fomo Live Feed 与 Frontrun")
    set_run_font(run, size=22, bold=True, color="000000")
    run.add_break()
    run = title.add_run("快捷交易合作接入说明")
    set_run_font(run, size=22, bold=True, color="000000")
    title_p_pr = title._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(3)
    run = subtitle.add_run("供 Frontrun 团队进行产品与技术可行性评估")
    set_run_font(run, size=12, color=MUTED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run("提案方  Fomo Live Feed 项目方    版本  1.0    日期  2026 年 9 月 9 日")
    set_run_font(run, size=9.5, color=MUTED)

    add_heading(doc, "结论与合作请求", 1)
    add_rich_paragraph(
        doc,
        [
            ("我们希望在 Fomo Live Feed 的实时交易信号卡片中接入 Frontrun 快捷交易能力。", True),
            ("Fomo Live Feed 负责提供经过校验的链、代币合约地址和上下文信号；Frontrun 负责钱包、交易参数、确认、签名和成交。用户点击一次即可进入 Frontrun 的 Quick Buy 或 Instant Trade 界面，并由用户最终确认交易。", False),
        ],
        after=7,
    )
    add_rich_paragraph(
        doc,
        [
            ("建议先完成 Solana 与 BNB Chain 的小范围 POC。", True),
            ("目标方案为浏览器扩展之间的版本化消息交接；若该能力尚未开放，可先用带入 chain 与 CA 的 Deep Link 验证用户需求，再升级为原位快捷交易。", False),
        ],
        after=7,
    )
    add_rich_paragraph(
        doc,
        [
            ("需要 Frontrun 团队确认并提供：", True),
            ("可接受的合作接入方式、支持的链与 CA 格式、调用方鉴权或白名单机制、打开快捷交易界面的协议，以及成功或失败的最小响应契约。", False),
        ],
        after=10,
    )

    add_heading(doc, "合作背景", 1)
    add_rich_paragraph(
        doc,
        [
            ("Fomo Live Feed ", True),
            ("是一款 Chrome 扩展，读取当前已登录 Fomo 用户所关注交易者的实时动态，并在 Chrome 侧边栏或始终置顶窗口中展示统一信息流。用户可查看交易者、动作、代币、链、金额、事件市值和 CA，并进行筛选、搜索、备注及安全跳转。", False),
        ],
    )
    add_rich_paragraph(
        doc,
        [
            ("Frontrun ", True),
            ("公开产品信息显示，其扩展已覆盖钱包标签、Twitter 与链上情报、交易钱包、Quick Buy、Instant Trade、交易预设及自动化功能。双方的能力天然互补：Fomo Live Feed 缩短“发现信号”的时间，Frontrun 缩短“从信号到安全交易确认”的时间。", False),
        ],
        after=9,
    )

    add_table(
        doc,
        ["环节", "Fomo Live Feed 当前能力", "Frontrun 可承接能力"],
        [
            ["发现", "实时捕获所关注交易者的买入、卖出及观点等事件", "无需替换现有信号源"],
            ["识别", "输出标准化 chain、tokenAddress、symbol 和事件上下文", "校验链与 CA，解析可交易资产"],
            ["决策", "在紧凑卡片中展示信号、金额、市值、交易者和时间", "加载用户自己的交易预设、钱包与风险参数"],
            ["执行", "当前版本不下单、不连接钱包、不读取私钥", "展示交易确认界面并负责签名、提交与结果反馈"],
        ],
        [0.75, 2.95, 2.95],
        font_size=9.2,
    )

    add_heading(doc, "当前产品边界", 1)
    add_bullet(doc, "现有生产版本为只读信息工具，不会下单、跟单或修改任何钱包状态。")
    add_bullet(doc, "扩展只接收经过运行时校验的事件字段，不导出 Fomo Cookie、请求头或认证令牌。")
    add_bullet(doc, "事件历史、筛选设置与用户备注保存在本地；默认保留 30 天或最多 20,000 条。")
    add_bullet(doc, "当前代码库尚未包含 Frontrun SDK、交易 API 或扩展间消息接入。官网公开的 Data API 主要提供 Twitter 与钱包情报，也未公开交易执行接口。")
    add_bullet(doc, "因此本说明是一份合作接入提案，不代表交易能力已经联调或上线。")

    add_heading(doc, "推荐接入方案", 1)
    add_heading(doc, "目标方案  扩展间快捷交易交接", 2)
    add_number(doc, 1, "用户在 Fomo Live Feed 卡片上点击 Frontrun 快捷买入。")
    add_number(doc, 2, "Fomo Live Feed 校验 chain 与 tokenAddress，生成短时有效、不可重放的请求。")
    add_number(doc, 3, "请求通过 Frontrun 认可的扩展间通信机制发送；接收方按合作方 ID、协议版本、字段白名单和 CA 规则再次校验。")
    add_number(doc, 4, "Frontrun 打开自己的 Quick Buy 或 Instant Trade 界面，并加载用户现有钱包、买入预设、滑点、优先费及风控设置。")
    add_number(doc, 5, "用户在 Frontrun 界面核对代币、链、金额和钱包后确认。Frontrun 独立完成签名、广播及交易结果处理。")
    add_number(doc, 6, "Fomo Live Feed 只接收 accepted、unsupported 或 invalid 等最小状态，不接触密钥、签名和认证材料。")

    add_heading(doc, "兼容方案  Deep Link", 2)
    add_rich_paragraph(
        doc,
        [
            ("若 Frontrun 暂未开放扩展间接口，", True),
            ("可先提供固定 HTTPS Deep Link 或受支持终端跳转规则。Fomo Live Feed 只拼接已校验的 chain 与 tokenAddress，打开 Frontrun 或其支持的交易页面，由 Frontrun 扩展识别页面 CA 并展示交易控件。该方案适合作为一到两周内的低成本 POC。", False),
        ],
        after=8,
    )

    add_heading(doc, "不推荐方案  我方直连交易后端", 2)
    add_rich_paragraph(
        doc,
        [
            ("首期不建议由 Fomo Live Feed 直接调用交易执行 API。", True),
            ("这会让我方承担钱包认证、交易参数、签名授权、错误恢复与安全审计责任，也会破坏当前只读扩展的信任边界。若未来确有自动化交易需求，应另立项目并单独完成权限、合规与风控评审。", False),
        ],
        after=8,
    )

    add_heading(doc, "建议接口契约", 1)
    add_rich_paragraph(
        doc,
        [
            ("下面字段用于表达合作意图，最终命名与传输机制以 Frontrun 的接口规范为准。", False),
            ("首期请求只要求打开并预填交易界面，不直接执行交易。", True),
        ],
        after=7,
    )
    add_table(
        doc,
        ["字段", "要求", "说明"],
        [
            ["version", "Required", "协议版本，例如 1"] ,
            ["requestId", "Required", "由调用方生成的唯一请求 ID，用于去重与排障"],
            ["source", "Required", "固定为 fomo-live-feed，接收方需做白名单校验"],
            ["intent", "Required", "首期固定为 open-quick-buy，不表示授权成交"],
            ["chain", "Required", "首期只允许 solana 或 bnb"],
            ["tokenAddress", "Required", "经链特定规则校验后的完整 CA"],
            ["tokenSymbol", "Optional", "仅用于展示，不作为资产识别依据"],
            ["occurredAt", "Optional", "原始信号发生时间，Unix 毫秒"],
            ["expiresAt", "Required", "请求过期时间，建议不超过生成后 30 秒"],
        ],
        [1.1, 1.0, 4.55],
        font_size=8.8,
    )

    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.28)
    code.paragraph_format.right_indent = Inches(0.2)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.05
    code_text = """{
  \"version\": 1,
  \"requestId\": \"01J...\",
  \"source\": \"fomo-live-feed\",
  \"intent\": \"open-quick-buy\",
  \"chain\": \"solana\",
  \"tokenAddress\": \"<validated contract address>\",
  \"tokenSymbol\": \"TOKEN\",
  \"occurredAt\": 1788940800000,
  \"expiresAt\": 1788940830000
}"""
    run = code.add_run(code_text)
    run.font.name = "Menlo"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Menlo")
    run.font.size = Pt(8.6)
    run.font.color.rgb = RGBColor.from_string(INK)

    add_heading(doc, "响应与错误处理", 2)
    add_bullet(doc, "accepted：Frontrun 已接受请求并打开交易界面，不表示交易成功。")
    add_bullet(doc, "unsupported：当前链、资产或客户端版本不支持。")
    add_bullet(doc, "invalid：请求过期、字段非法、调用方未授权或 CA 校验失败。")
    add_bullet(doc, "用户取消、余额不足、报价变化和链上失败均由 Frontrun 自己的交易流程处理；我方不推断成交结果。")

    add_heading(doc, "安全与隐私边界", 1)
    add_table(
        doc,
        ["原则", "约束"],
        [
            ["最小数据", "只传 chain、CA、请求 ID、意图和必要展示信息；不传 Cookie、Authorization、私钥、助记词或签名。"],
            ["双重校验", "发送方和接收方都执行链特定 CA 校验；tokenSymbol 不能用于资产定位。"],
            ["显式确认", "首期每次交易都必须在 Frontrun 界面由用户确认，不启用自动买入或静默成交。"],
            ["调用方限制", "Frontrun 应按扩展 ID 或合作方标识设置白名单，并校验协议版本、有效期和重复 requestId。"],
            ["职责隔离", "Frontrun 独立管理钱包、交易预设、报价、签名、广播及失败恢复；Fomo Live Feed 保持信息发现工具定位。"],
            ["可观测性", "双方日志仅记录请求 ID、版本、链、状态码和时间，不记录认证材料或完整交易凭据。"],
        ],
        [1.25, 5.4],
        font_size=9.1,
    )

    add_heading(doc, "首期 POC 范围", 1)
    add_bullet(doc, "链范围：Solana 与 BNB Chain。")
    add_bullet(doc, "入口范围：Fomo Live Feed 侧边栏与始终置顶窗口中的买入事件卡片。")
    add_bullet(doc, "行为范围：用户点击后打开 Frontrun 交易确认界面，并预填正确链与 CA。")
    add_bullet(doc, "预设范围：买入金额、钱包、滑点与费用继续采用用户在 Frontrun 中已有的设置。")
    add_bullet(doc, "明确排除：自动跟单、无确认交易、卖出、限价单、TP/SL、批量交易、历史回放触发及交易结果写回。")

    add_heading(doc, "POC 验收标准", 2)
    add_number(doc, 1, "十个受支持测试代币均能从卡片打开 Frontrun，并显示与原事件一致的链和完整 CA。")
    add_number(doc, 2, "非法、缺失、过期或链不匹配的请求全部被拒绝，且不会打开可确认的交易表单。")
    add_number(doc, 3, "重复点击不会产生多个交易请求或意外成交。")
    add_number(doc, 4, "用户取消不会改变钱包状态，也不会在 Fomo Live Feed 中显示为已成交。")
    add_number(doc, 5, "整个链路中 Fomo Live Feed 无法读取 Frontrun 私钥、助记词、签名材料或认证令牌。")
    add_number(doc, 6, "双方能以 requestId 定位一次失败交接，同时日志不包含敏感信息。")

    add_heading(doc, "双方工作边界", 1)
    add_table(
        doc,
        ["事项", "Fomo Live Feed", "Frontrun"],
        [
            ["信号与 UI", "提供事件卡片、点击入口和必要上下文", "提供按钮文案、品牌资源与交互规范"],
            ["协议", "按约定生成、校验、过期和去重请求", "提供接收端、白名单、版本与响应规范"],
            ["资产校验", "发送前校验 chain 与 CA", "接收后再次校验并确认资产可交易"],
            ["钱包与交易", "不接触钱包和签名", "负责钱包、报价、预设、确认、签名及广播"],
            ["测试", "提供测试事件、测试扩展和请求日志", "提供测试环境或安全测试钱包与状态日志"],
        ],
        [1.1, 2.75, 2.8],
        font_size=9.0,
    )

    add_heading(doc, "请 Frontrun 团队确认", 1)
    add_number(doc, 1, "是否愿意支持第三方 Chrome 扩展调用 Frontrun 的快捷交易界面。")
    add_number(doc, 2, "推荐采用扩展间消息、公开 Deep Link、Partner SDK，还是在指定 HTTPS 页面注入交易控件。")
    add_number(doc, 3, "Solana 与 BNB Chain 的正式 chain 标识、CA 规则以及当前可交易资产限制。")
    add_number(doc, 4, "能否只打开并预填交易界面，由 Frontrun 保持完整钱包与签名边界。")
    add_number(doc, 5, "调用方白名单、测试环境、版本兼容、限流、超时和错误码要求。")
    add_number(doc, 6, "是否允许使用 Frontrun 名称、Logo 和 Quick Buy 文案，以及相应品牌规范。")
    add_number(doc, 7, "若合作进入正式版本，双方对支持、监控、安全事件和接口变更的联系人与响应方式。")

    add_heading(doc, "建议推进步骤", 1)
    add_number(doc, 1, "Frontrun 选择首选接入方式并回复上述确认项。")
    add_number(doc, 2, "双方用一组固定测试 CA 完成最小协议联调，不接入真实资金。")
    add_number(doc, 3, "在测试钱包中验证用户确认、取消、非法请求、重复请求及不支持链。")
    add_number(doc, 4, "POC 验收后再确定正式发布范围、品牌展示、接口稳定性承诺与联合运营方式。")

    add_heading(doc, "资料来源与说明", 1)
    add_table(
        doc,
        ["资料", "链接"],
        [
            ["Frontrun 官网", "https://www.frontrun.pro/zh"],
            ["Frontrun Data API", "https://www.frontrun.pro/zh/data-api"],
            ["Frontrun 功能指南", "https://www.frontrun.pro/zh/how-to"],
            ["Frontrun 隐私政策", "https://www.frontrun.pro/zh/privacy"],
        ],
        [1.75, 4.9],
        font_size=9.2,
    )
    add_rich_paragraph(
        doc,
        [
            ("说明：", True),
            ("Frontrun 的公开资料证明产品具备快捷交易、钱包和数据能力，但未公开第三方交易执行接口。本说明中的消息结构、状态码和 POC 安排均为合作建议，需经 Frontrun 团队确认后方可作为正式接口依据。", False),
        ],
        after=4,
    )

    return doc


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
